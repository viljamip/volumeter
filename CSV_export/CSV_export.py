import pandas as pd
import numpy as np


# requires openpyxl
data = pd.read_pickle('sampleData2')

def toExcel(df, name):
	from openpyxl import Workbook
	from openpyxl.chart import ScatterChart, Reference, Series
	from openpyxl import load_workbook
	df = df.dropna(axis='columns')
	df.index.name = 'Depth'
	df.to_excel('measurements/{}.xlsx'.format(name))
	workbook = load_workbook(filename='measurements/{}.xlsx'.format(name))
	sheet = workbook.active
	chart = ScatterChart()
	chart.title = 'Measurement'
	chart.style = 1
	chart.x_axis.title = 'Depth (mm)'
	chart.y_axis.title = 'Volume (ml)'

	num_rows = len(sheet['A'])
	xvalues = Reference(sheet, min_col=1, min_row=1, max_row=num_rows)
	values = Reference(sheet, min_col=2, min_row=1, max_row=num_rows)
	series = Series(values, xvalues, title_from_data=True)
	chart.series.append(series)
	chart.legend = None
	sheet.add_chart(chart, "C1")

	workbook.save('measurements/{}.xlsx'.format(name))

toExcel(data, 'testingFunc')
